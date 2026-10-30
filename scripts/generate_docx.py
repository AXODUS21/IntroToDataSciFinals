"""
Converts FINAL_PROJECT_DOCUMENTATION.md into a professionally formatted Microsoft Word (.docx)
technical report with native Word OMML equations, formatted inline math, syntax-styled code/output
blocks, styled academic tables, clickable hyperlinks, and embedded high-resolution figures.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls


# Color Palette
GREEN_DARK = "1B5E20"      # Deep Lasallian Green for table headers & accents
GREEN_ACCENT = "2E7D32"    # Medium Green for borders
ROW_ALT_FILL = "F4F7F4"    # Soft tinted alternating table row
ROW_WHITE_FILL = "FFFFFF"  # Explicit white fill for even table rows
CODE_BG_FILL = "F6F8FA"    # Subtle gray fill for programmatic output blocks
CODE_BORDER = "D0D7DE"     # Subtle gray border for code blocks

RGB_GREEN_HEADER = RGBColor(27, 94, 32)
RGB_GREEN_SUB = RGBColor(46, 125, 50)
RGB_DARK_TEXT = RGBColor(30, 34, 39)
RGB_MUTED_TEXT = RGBColor(90, 98, 106)
RGB_WHITE = RGBColor(255, 255, 255)
RGB_LINK = RGBColor(37, 99, 235)


def set_cell_background(cell, hex_color):
    """Sets background shading of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:shd")):
        tcPr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets internal padding (in dxa) for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Sets custom borders on a cell. Each side is a tuple (val, sz, color)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side, spec in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if spec:
            val, sz, color = spec
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), val)
            b.set(qn("w:sz"), str(sz))
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), color)
            tcBorders.append(b)
    tcPr.append(tcBorders)


def set_row_cant_split(row):
    """Prevents a table row from breaking across a page boundary."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def set_repeat_table_header(row):
    """Marks a table row as a repeating header across page breaks."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))


def add_paragraph_bottom_border(paragraph, color_hex=GREEN_ACCENT, sz="12"):
    """Adds a clean bottom accent border under a heading paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_hyperlink(paragraph, text, url, font_size=Pt(10), force_color=None):
    """Adds a clickable hyperlink run to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)

    c = OxmlElement("w:color")
    c.set(qn("w:val"), "2563EB")
    rPr.append(c)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    if font_size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(font_size.pt * 2)))
        rPr.append(sz)

    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_styled_run(
    paragraph,
    text,
    bold=False,
    italic=False,
    code=False,
    subscript=False,
    superscript=False,
    font_name=None,
    font_size=Pt(10.5),
    font_color=None,
):
    """Appends a run with the specified typographic properties."""
    if not text:
        return None
    r = paragraph.add_run(text)
    if code:
        r.font.name = "Consolas"
        r.font.size = Pt(max(8.5, font_size.pt - 1.0)) if font_size else Pt(9.5)
    else:
        r.font.name = font_name or "Calibri"
        if font_size:
            r.font.size = font_size
    if bold:
        r.font.bold = True
    if italic:
        r.font.italic = True
    if subscript:
        r.font.subscript = True
    if superscript:
        r.font.superscript = True
    if font_color is not None:
        r.font.color.rgb = font_color
    return r


def render_inline_math(paragraph, expr, base_bold=False, base_italic=False, font_size=Pt(10.5), font_color=None):
    """
    Converts an inline LaTeX expression (without surrounding $) into properly formatted Word runs
    using Unicode mathematical symbols, real subscripts/superscripts, and italic math variables.
    """
    s = expr.strip()

    # Special exact matches for complex mathematical expressions in this document
    if s == r"\mathbf{x}_t \in \mathbb{R}^{d}":
        add_styled_run(paragraph, "x", bold=True, italic=True, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, "t", italic=True, subscript=True, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, " ∈ ℝ", bold=base_bold, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, "d", italic=True, superscript=True, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        return

    if s == r"y_t \in \{0, 1\}":
        add_styled_run(paragraph, "y", bold=base_bold, italic=True, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, "t", bold=base_bold, italic=True, subscript=True, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, " ∈ {0, 1}", bold=base_bold, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        return

    if s == r"\in \{0, 1\}":
        add_styled_run(paragraph, "∈ {0, 1}", bold=base_bold, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        return

    if s == r"\text{MJ/m}^2":
        add_styled_run(paragraph, "MJ/m", bold=base_bold, italic=base_italic, font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, "2", bold=base_bold, superscript=True, font_size=font_size, font_color=font_color)
        return

    if s == r"\text{AUC} = \mathbf{0.966}":
        add_styled_run(paragraph, "AUC = ", bold=base_bold, italic=base_italic, font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, "0.966", bold=True, italic=base_italic, font_size=font_size, font_color=font_color)
        return

    # General replacements for \text{...}, \mathbf{...}, symbols
    s = re.sub(r"\\text\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\mathbf\{([^}]*)\}", r"\1", s)
    s = s.replace(r"\pm", "±")
    s = s.replace(r"\ge", "≥")
    s = s.replace(r"\le", "≤")
    s = s.replace(r"\times", "×")
    s = s.replace(r"\in", "∈")
    s = s.replace(r"\{", "{")
    s = s.replace(r"\}", "}")

    # Single-letter mathematical variables (t, r, X, y, N) or "r = ...", "N = ..."
    m_var_eq = re.match(r"^([rXyNt])\s*=\s*(.+)$", s)
    if m_var_eq:
        var_name, rest = m_var_eq.group(1), m_var_eq.group(2)
        add_styled_run(paragraph, var_name, bold=base_bold, italic=True, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        add_styled_run(paragraph, f" = {rest}", bold=base_bold, italic=base_italic, font_size=font_size, font_color=font_color)
        return

    if s in ("t", "r", "X", "y", "N"):
        add_styled_run(paragraph, s, bold=base_bold, italic=True, font_name="Cambria Math", font_size=font_size, font_color=font_color)
        return

    add_styled_run(paragraph, s, bold=base_bold, italic=base_italic, font_size=font_size, font_color=font_color)


# Token pattern for inline formatting:
# 1. Markdown links: [text](url)
# 2. Bold: **...**
# 3. Italic: *...*
# 4. Inline code: `...`
# 5. Inline math: $...$
INLINE_TOKEN_RE = re.compile(
    r"(\[[^\]]+\]\([^)]+\))"        # [label](url)
    r"|(\*\*(?:[^*]|\*(?!\*))+\*\*)" # **bold**
    r"|(\*(?:[^*]+)\*)"              # *italic*
    r"|(`[^`]+`)"                    # `code`
    r"|(\$[^$]+\$)"                  # $math$
)


def add_formatted_text(
    paragraph,
    text,
    base_bold=False,
    base_italic=False,
    font_size=Pt(10.5),
    font_color=None,
):
    """
    Recursively parses markdown inline syntax (**bold**, *italic*, `code`, $math$, [link](url))
    and appends properly formatted runs to `paragraph`.
    """
    if not text:
        return

    pos = 0
    for match in INLINE_TOKEN_RE.finditer(text):
        start, end = match.span()
        if start > pos:
            plain = text[pos:start]
            add_styled_run(
                paragraph,
                plain,
                bold=base_bold,
                italic=base_italic,
                font_size=font_size,
                font_color=font_color,
            )

        link_tok, bold_tok, italic_tok, code_tok, math_tok = match.groups()

        if link_tok:
            m_link = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", link_tok)
            if m_link:
                label, url = m_link.group(1), m_link.group(2)
                add_hyperlink(paragraph, label, url, font_size=font_size, force_color=font_color)
            else:
                add_styled_run(paragraph, link_tok, bold=base_bold, italic=base_italic, font_size=font_size, font_color=font_color)

        elif bold_tok:
            inner = bold_tok[2:-2]
            add_formatted_text(
                paragraph,
                inner,
                base_bold=True,
                base_italic=base_italic,
                font_size=font_size,
                font_color=font_color,
            )

        elif italic_tok:
            inner = italic_tok[1:-1]
            add_formatted_text(
                paragraph,
                inner,
                base_bold=base_bold,
                base_italic=True,
                font_size=font_size,
                font_color=font_color,
            )

        elif code_tok:
            inner = code_tok[1:-1]
            add_styled_run(
                paragraph,
                inner,
                bold=base_bold,
                italic=base_italic,
                code=True,
                font_size=font_size,
                font_color=font_color,
            )

        elif math_tok:
            inner = math_tok[1:-1]
            render_inline_math(
                paragraph,
                inner,
                base_bold=base_bold,
                base_italic=base_italic,
                font_size=font_size,
                font_color=font_color,
            )

        pos = end

    if pos < len(text):
        add_styled_run(
            paragraph,
            text[pos:],
            bold=base_bold,
            italic=base_italic,
            font_size=font_size,
            font_color=font_color,
        )


# ---------------------------------------------------------------------------
# Native Word OMML Display Equation Builder
# ---------------------------------------------------------------------------

def _omml_r(text, italic=False, bold=False):
    """Creates an <m:r> element string for OMML equations."""
    nor_xml = "<m:nor/>" if not italic else ""
    sty_val = "b" if (bold and not italic) else ("bi" if (bold and italic) else ("i" if italic else "p"))
    sty_xml = f'<m:sty m:val="{sty_val}"/>'
    # Escape XML special characters
    esc = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    return (
        f"<m:r>"
        f"<m:rPr>{nor_xml}{sty_xml}</m:rPr>"
        f'<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr>'
        f'<m:t xml:space="preserve">{esc}</m:t>'
        f"</m:r>"
    )


def _omml_sub(base_xml, sub_xml):
    return f"<m:sSub><m:e>{base_xml}</m:e><m:sub>{sub_xml}</m:sub></m:sSub>"


def _omml_frac(num_xml, den_xml):
    return f"<m:f><m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>"


def _omml_delim(inner_xml, beg="(", end=")"):
    return (
        f"<m:d>"
        f'<m:dPr><m:begChr m:val="{beg}"/><m:endChr m:val="{end}"/></m:dPr>'
        f"<m:e>{inner_xml}</m:e>"
        f"</m:d>"
    )


def build_omml_for_display_latex(latex_str):
    """
    Returns an <m:oMathPara> lxml element for each display equation in the report.
    """
    s = " ".join(latex_str.strip().split())

    if r"\begin{cases}" in s:
        # Target variable piecewise definition
        row1 = (
            _omml_r("0,    if Rented Bike Count", italic=False)
            + _omml_sub(_omml_r("", italic=False), _omml_r("t", italic=True))
            + _omml_r(" = 0   (No bikes rented / Zero demand)", italic=False)
        )
        # Better: attach subscript t directly to Count
        row1 = (
            _omml_r("0,   if ", italic=False)
            + _omml_sub(_omml_r("Rented Bike Count", italic=False), _omml_r("t", italic=True))
            + _omml_r(" = 0   (No bikes rented / Zero demand)", italic=False)
        )
        row2 = (
            _omml_r("1,   if ", italic=False)
            + _omml_sub(_omml_r("Rented Bike Count", italic=False), _omml_r("t", italic=True))
            + _omml_r(" > 0   (At least one bike rented / Active demand)", italic=False)
        )
        cases_xml = (
            f"<m:d>"
            f'<m:dPr><m:begChr m:val="{{"/><m:endChr m:val=""/></m:dPr>'
            f"<m:e><m:eqArr><m:e>{row1}</m:e><m:e>{row2}</m:e></m:eqArr></m:e>"
            f"</m:d>"
        )
        body_xml = _omml_sub(_omml_r("y", italic=True), _omml_r("t", italic=True)) + _omml_r(" = ", italic=False) + cases_xml

    elif r"\mu_{\text{train}}" in s:
        # z = (x - \mu_train) / \sigma_train
        num = _omml_r("x", italic=True) + _omml_r(" − ", italic=False) + _omml_sub(_omml_r("μ", italic=True), _omml_r("train", italic=False))
        den = _omml_sub(_omml_r("σ", italic=True), _omml_r("train", italic=False))
        body_xml = _omml_r("z", italic=True) + _omml_r(" = ", italic=False) + _omml_frac(num, den)

    elif r"2 \times N_j" in s:
        # w_j = N / (2 * N_j)
        num = _omml_r("N", italic=True)
        den = _omml_r("2 × ", italic=False) + _omml_sub(_omml_r("N", italic=True), _omml_r("j", italic=True))
        body_xml = _omml_sub(_omml_r("w", italic=True), _omml_r("j", italic=True)) + _omml_r(" = ", italic=False) + _omml_frac(num, den)

    elif s.startswith(r"\text{Accuracy}"):
        body_xml = (
            _omml_r("Accuracy = ", italic=False)
            + _omml_frac(
                _omml_r("TP + TN", italic=False),
                _omml_r("TP + TN + FP + FN", italic=False),
            )
        )

    elif s.startswith(r"\text{Balanced Accuracy}"):
        frac1 = _omml_frac(_omml_r("Sensitivity + Specificity", italic=False), _omml_r("2", italic=False))
        half = _omml_frac(_omml_r("1", italic=False), _omml_r("2", italic=False))
        inner_sum = (
            _omml_frac(_omml_r("TP", italic=False), _omml_r("TP + FN", italic=False))
            + _omml_r(" + ", italic=False)
            + _omml_frac(_omml_r("TN", italic=False), _omml_r("TN + FP", italic=False))
        )
        body_xml = (
            _omml_r("Balanced Accuracy = ", italic=False)
            + frac1
            + _omml_r(" = ", italic=False)
            + half
            + _omml_delim(inner_sum, "(", ")")
        )

    elif s.startswith(r"\text{Precision}_c"):
        lhs = _omml_sub(_omml_r("Precision", italic=False), _omml_r("c", italic=True))
        num = _omml_sub(_omml_r("True Positives", italic=False), _omml_r("c", italic=True))
        den = (
            _omml_sub(_omml_r("True Positives", italic=False), _omml_r("c", italic=True))
            + _omml_r(" + ", italic=False)
            + _omml_sub(_omml_r("False Positives", italic=False), _omml_r("c", italic=True))
        )
        body_xml = lhs + _omml_r(" = ", italic=False) + _omml_frac(num, den)

    elif s.startswith(r"\text{Recall}_1"):
        lhs = _omml_sub(_omml_r("Recall", italic=False), _omml_r("1", italic=False))
        body_xml = lhs + _omml_r(" = ", italic=False) + _omml_frac(_omml_r("TP", italic=False), _omml_r("TP + FN", italic=False))

    elif s.startswith(r"\text{Specificity}"):
        lhs = _omml_r("Specificity = ", italic=False) + _omml_sub(_omml_r("Recall", italic=False), _omml_r("0", italic=False))
        body_xml = lhs + _omml_r(" = ", italic=False) + _omml_frac(_omml_r("TN", italic=False), _omml_r("TN + FP", italic=False))

    elif r"F_{1, \text{macro}}" in s:
        lhs = _omml_sub(_omml_r("F", italic=True), _omml_r("1, macro", italic=False))
        f1_c0 = _omml_sub(_omml_r("F", italic=True), _omml_r("1, Class 0", italic=False))
        f1_c1 = _omml_sub(_omml_r("F", italic=True), _omml_r("1, Class 1", italic=False))
        frac_macro = _omml_frac(f1_c0 + _omml_r(" + ", italic=False) + f1_c1, _omml_r("2", italic=False))
        f1_sym = _omml_sub(_omml_r("F", italic=True), _omml_r("1", italic=False))
        frac_f1 = _omml_frac(_omml_r("Precision × Recall", italic=False), _omml_r("Precision + Recall", italic=False))
        body_xml = (
            lhs
            + _omml_r(" = ", italic=False)
            + frac_macro
            + _omml_r(" ,    where  ", italic=False)
            + f1_sym
            + _omml_r(" = 2 × ", italic=False)
            + frac_f1
        )
    else:
        body_xml = _omml_r(s, italic=False)

    xml_str = f"<m:oMathPara {nsdecls('m', 'w')}><m:oMath>{body_xml}</m:oMath></m:oMathPara>"
    return parse_xml(xml_str)


def add_display_equation(doc, latex_str, indent_inches=0.0):
    """Adds a centered native Word OMML display equation paragraph."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    if indent_inches > 0:
        p.paragraph_format.left_indent = Inches(indent_inches)
    omml_el = build_omml_for_display_latex(latex_str)
    p._p.append(omml_el)


def add_code_block(doc, code_lines):
    """Renders a programmatic output / code block inside a styled, shaded monospace container."""
    if not code_lines:
        return
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, CODE_BG_FILL)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=140)
    set_cell_borders(
        cell,
        top=("single", 4, CODE_BORDER),
        bottom=("single", 4, CODE_BORDER),
        right=("single", 4, CODE_BORDER),
        left=("single", 18, GREEN_ACCENT),
    )

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1

    text_content = "\n".join(code_lines)
    r = p.add_run(text_content)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGB_DARK_TEXT

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(3)


def configure_headers_and_footers(doc):
    """Adds a subtle running header and footer to the document."""
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.95)
        section.right_margin = Inches(0.95)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Intdsci Final Project Report — Section C2B (Group FML) | De La Salle Lipa")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGB_MUTED_TEXT

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Seoul Bike Sharing Demand Binary Classification — Fabello, Lubi, & Macaraig")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGB_MUTED_TEXT


def build_cover_header(doc):
    """Builds a formal academic header and structured metadata summary table."""
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(4)

    r1 = p_inst.add_run("DE LA SALLE LIPA\n")
    r1.font.name = "Calibri"
    r1.font.size = Pt(15)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(22, 138, 74)

    r2 = p_inst.add_run("College of Information Technology and Engineering\nComputer Science Department")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10.5)
    r2.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(10)
    add_paragraph_bottom_border(p_title, color_hex=GREEN_ACCENT, sz="12")

    rt1 = p_title.add_run("COMPREHENSIVE PROJECT DOCUMENTATION\n")
    rt1.font.name = "Calibri"
    rt1.font.size = Pt(16)
    rt1.font.bold = True

    rt2 = p_title.add_run(
        "Supervised Binary Classification of Hourly Bike Rental Demand in Seoul\n"
        "Using Ambient Meteorological and Calendar Features"
    )
    rt2.font.name = "Calibri"
    rt2.font.size = Pt(11.5)
    rt2.font.bold = True
    rt2.font.italic = True

    # Metadata summary table
    meta_rows = [
        ("Course & Section", "**Introduction to Data Science (Intdsci)** — Section **C2B**"),
        ("Course Instructor", "**Ms. Donna Garcia**"),
        ("Group Name & Members", "**Group FML:** Fabello, Courtney Judge; Lubi, Axellerosh; Macaraig, Paolo Adrian"),
        ("Date of Submission", "Friday, October 30, 2026"),
        ("Primary Dataset", "Seoul Bike Sharing Demand Dataset (UCI Machine Learning Repository & Class Google Drive)"),
        (
            "Project Deliverables",
            "Executable Python Pipeline (`run_pipeline.py`), Reproducible Jupyter Notebook (`notebooks/final_project.ipynb`), and Formatted Project Report (`reports/FINAL_PROJECT_DOCUMENTATION.docx`)",
        ),
    ]

    t = doc.add_table(rows=len(meta_rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"

    col_widths = [Inches(1.85), Inches(4.75)]
    for r_idx, (label, val) in enumerate(meta_rows):
        row = t.rows[r_idx]
        set_row_cant_split(row)

        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = col_widths[0]
        c1.width = col_widths[1]

        set_cell_background(c0, GREEN_DARK)
        set_cell_margins(c0, top=70, bottom=70, left=120, right=120)
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        add_styled_run(p0, label, bold=True, font_size=Pt(9.5), font_color=RGB_WHITE)

        bg = ROW_ALT_FILL if r_idx % 2 == 1 else ROW_WHITE_FILL
        set_cell_background(c1, bg)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        add_formatted_text(p1, val, font_size=Pt(9.5), font_color=RGB_DARK_TEXT)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(4)
    spacer.paragraph_format.space_after = Pt(4)


def create_word_report():
    doc = Document()

    # Configure default Normal style font
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)

    configure_headers_and_footers(doc)
    build_cover_header(doc)

    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    md_path = os.path.join(base_dir, "reports", "FINAL_PROJECT_DOCUMENTATION.md")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_rows = []
    table_alignments = []

    in_code_block = False
    code_lines = []

    in_math_block = False
    math_lines = []

    def flush_table():
        nonlocal in_table, table_rows, table_alignments
        if not table_rows:
            in_table = False
            return

        num_cols = len(table_rows[0])
        t = doc.add_table(rows=len(table_rows), cols=num_cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"

        for r_idx, row_vals in enumerate(table_rows):
            row = t.rows[r_idx]
            set_row_cant_split(row)
            if r_idx == 0:
                set_repeat_table_header(row)

            for c_idx in range(min(num_cols, len(row_vals))):
                cell_value = row_vals[c_idx].strip()
                cell = row.cells[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=75, bottom=75, left=100, right=100)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.08

                # Apply column alignment from markdown separator row
                if c_idx < len(table_alignments):
                    align = table_alignments[c_idx]
                    if align == "center":
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif align == "right":
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif c_idx > 0 and num_cols >= 4:
                    # Default center numeric columns if unspecified
                    if re.match(r"^[\d$+\-±.%(),\s*]+$", cell_value) and len(cell_value) < 28:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                if r_idx == 0:
                    set_cell_background(cell, GREEN_DARK)
                    add_formatted_text(
                        p,
                        cell_value,
                        base_bold=True,
                        font_size=Pt(9.5),
                        font_color=RGB_WHITE,
                    )
                else:
                    bg = ROW_ALT_FILL if r_idx % 2 == 1 else ROW_WHITE_FILL
                    set_cell_background(cell, bg)
                    add_formatted_text(
                        p,
                        cell_value,
                        base_bold=False,
                        font_size=Pt(9.0),
                        font_color=RGB_DARK_TEXT,
                    )

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(2)
        spacer.paragraph_format.space_after = Pt(4)
        table_rows = []
        table_alignments = []
        in_table = False

    # Start parsing after the top metadata block (which ends at line 19 '---')
    start_idx = 0
    for idx, line in enumerate(lines):
        if line.strip() == "## Executive Summary":
            start_idx = idx
            break

    for raw_line in lines[start_idx:]:
        line_no_nl = raw_line.rstrip("\r\n")
        stripped = line_no_nl.strip()

        # 1. Handle fenced code / programmatic output blocks
        if stripped.startswith("```"):
            if in_table:
                flush_table()
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code_block = False
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line_no_nl)
            continue

        # 2. Handle multi-line or single-line display math blocks ($$ ... $$)
        if stripped.startswith("$$"):
            if in_table:
                flush_table()
            # Single-line $$ ... $$
            if stripped.endswith("$$") and len(stripped) > 4:
                latex_inner = stripped[2:-2].strip()
                lead_spaces = len(line_no_nl) - len(line_no_nl.lstrip(" "))
                add_display_equation(doc, latex_inner, indent_inches=0.25 if lead_spaces >= 2 else 0.0)
                continue
            else:
                if not in_math_block:
                    in_math_block = True
                    math_lines = []
                    rem = stripped[2:].strip()
                    if rem:
                        math_lines.append(rem)
                else:
                    rem = stripped[:-2].strip() if stripped.endswith("$$") else stripped
                    if rem:
                        math_lines.append(rem)
                    add_display_equation(doc, "\n".join(math_lines))
                    in_math_block = False
                    math_lines = []
                continue

        if in_math_block:
            if stripped.endswith("$$"):
                rem = stripped[:-2].strip()
                if rem:
                    math_lines.append(rem)
                add_display_equation(doc, "\n".join(math_lines))
                in_math_block = False
                math_lines = []
            else:
                math_lines.append(stripped)
            continue

        # 3. Handle Markdown tables
        if stripped.startswith("|") and stripped.endswith("|"):
            parts = [p.strip() for p in stripped[1:-1].split("|")]
            # Check if separator line like | :--- | :---: |
            if all(set(p).issubset({"-", ":", " "}) for p in parts if p):
                table_alignments = []
                for p in parts:
                    p_clean = p.strip()
                    if p_clean.startswith(":") and p_clean.endswith(":"):
                        table_alignments.append("center")
                    elif p_clean.endswith(":"):
                        table_alignments.append("right")
                    else:
                        table_alignments.append("left")
                continue
            table_rows.append(parts)
            in_table = True
            continue
        elif in_table:
            flush_table()

        # 4. Skip empty lines or raw markdown horizontal rules
        if not stripped or stripped == "---":
            continue

        # 5. Handle Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[2:].strip(), base_bold=True, font_size=Pt(15), font_color=RGBColor(22, 138, 74))
            add_paragraph_bottom_border(h, color_hex=GREEN_ACCENT, sz="10")

        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[3:].strip(), base_bold=True, font_size=Pt(13.5), font_color=RGBColor(22, 138, 74))
            add_paragraph_bottom_border(h, color_hex=GREEN_ACCENT, sz="8")

        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[4:].strip(), base_bold=True, font_size=Pt(11.5))

        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            add_formatted_text(h, stripped[5:].strip(), base_bold=True, font_size=Pt(10.5))

        # 6. Handle Embedded Images
        elif stripped.startswith("!["):
            try:
                caption = stripped[2 : stripped.index("]")]
                img_url = stripped[stripped.index("(") + 1 : stripped.index(")")]
                img_path = img_url.replace("file://", "").strip()
                if not os.path.exists(img_path):
                    cand_reports = os.path.normpath(os.path.join(base_dir, "reports", img_path))
                    cand_base = os.path.normpath(os.path.join(base_dir, img_path.lstrip("/\\")))
                    if os.path.exists(cand_reports):
                        img_path = cand_reports
                    elif os.path.exists(cand_base):
                        img_path = cand_base
                if os.path.exists(img_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(6)
                    p_img.paragraph_format.space_after = Pt(2)
                    p_img.paragraph_format.keep_with_next = True
                    p_img.add_run().add_picture(img_path, width=Inches(5.8))

                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(8)
                    add_formatted_text(p_cap, caption, base_bold=True, base_italic=True, font_size=Pt(9.5))
            except Exception as e:
                print(f"Image embed error: {e}")

        # 7. Handle Bullet Lists (including indented sub-bullets)
        elif stripped.startswith("* ") or stripped.startswith("- "):
            lead_spaces = len(line_no_nl) - len(line_no_nl.lstrip(" "))
            content = stripped[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            if lead_spaces >= 2:
                p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, content, font_size=Pt(10.5))

        # 8. Handle Numbered Lists (preserving exact markdown item numbers across equations/sub-bullets)
        elif re.match(r"^\d+\.\s+", stripped):
            m_num = re.match(r"^(\d+\.)\s+(.*)$", stripped)
            num_prefix, content = m_num.group(1), m_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            add_styled_run(p, f"{num_prefix} ", bold=True, font_size=Pt(10.5))
            add_formatted_text(p, content, font_size=Pt(10.5))

        # 9. Handle Indented Continuation Paragraphs (e.g., under numbered list or bullet)
        else:
            lead_spaces = len(line_no_nl) - len(line_no_nl.lstrip(" "))
            p = doc.add_paragraph()
            if lead_spaces >= 2:
                p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, stripped, font_size=Pt(10.5))

    if in_table:
        flush_table()
    if in_code_block and code_lines:
        add_code_block(doc, code_lines)

    output_docx = os.path.join(base_dir, "reports", "FINAL_PROJECT_DOCUMENTATION.docx")
    doc.save(output_docx)
    print(f"Successfully generated Word document at {output_docx}!")


if __name__ == "__main__":
    create_word_report()
